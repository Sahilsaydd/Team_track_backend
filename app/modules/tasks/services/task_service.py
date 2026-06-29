
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, update, desc
from sqlalchemy.orm import aliased
from fastapi import HTTPException, status
import os

from datetime import datetime, timedelta, timezone

from fastapi import HTTPException
from fastapi.responses import FileResponse

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession
from fastapi.responses import FileResponse
from app.modules.groups.models.group_member import GroupMember
from app.modules.tasks.models.task import Task
from app.modules.tasks.models.task_review import TaskReview
from app.modules.tasks.models.task_log import TaskLog
from app.modules.tasks.models.task_evidence import TaskEvidence
import os
from datetime import datetime, timedelta, timezone
from app.modules.notification.model.notification import (
    Notification
)
from sqlalchemy.orm import selectinload
from app.modules.users.models.user import User
from app.core.file_upload import save_upload_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from app.modules.tasks.enums import (
    ApprovalStatus,
    TaskStatus
)



async def create_task_service(db: AsyncSession,data,current_user):


    print("Current User id : ",current_user.id)
    print("Group Id :",data.group_id)

    leader_result = await db.execute(
        select(GroupMember).where(
            GroupMember.group_id == data.group_id,
            GroupMember.user_id == current_user.id,
            GroupMember.role_in_group == "Lead",
            GroupMember.is_active == True
        )
    )

    leader = leader_result.scalar_one_or_none()

    if not leader:

        raise HTTPException(
            status_code=403,
            detail="Only group leader can assign tasks"
        )


    member_result = await db.execute(
        select(GroupMember).where(
            GroupMember.group_id == data.group_id,
            GroupMember.user_id == data.assigned_to,
            GroupMember.is_active == True
        )
    )

    member = member_result.scalar_one_or_none()

    if not member:

        raise HTTPException(
            status_code=404,
            detail="Assigned user is not member of this group"
        )


    if member.role_in_group == "Lead":

        raise HTTPException(
            status_code=400,
            detail="Cannot assign task to another lead"
        )


    task = Task(

        title=data.title,

        description=data.description,

        priority=data.priority,
        task_type="Group",
        deadline=data.deadline,

        group_id=data.group_id,

        created_by=current_user.id,

        assigned_by=current_user.id,

        assigned_to=data.assigned_to,

        submitted_to=current_user.id,

        is_self_task=False
    )

    db.add(task)


    notification = Notification(

        user_id=data.assigned_to,

        title="New Group Task Assigned",

        message=f"You received task '{task.title}'",

        notification_type="task_assigned"
    )

    db.add(notification)

    await db.commit()

    await db.refresh(task)

    return task





async def create_hierarchy_task_service(db: AsyncSession,data,current_user):

    current_user_result = await db.execute(
        select(User)
        .options(selectinload(User.role))
        .where(User.id == current_user.id)
    )

    current_user = current_user_result.scalar_one_or_none()

    if not current_user:
        raise HTTPException(
            status_code=404,
            detail="Current user not found"
        )

    assigned_user_result = await db.execute(
        select(User)
        .options(selectinload(User.role))
        .where(User.id == data.assigned_to)
    )

    assigned_user = assigned_user_result.scalar_one_or_none()

    if not assigned_user:
        raise HTTPException(
            status_code=404,
            detail="Assigned user not found"
        )

    current_role = getattr(getattr(current_user, "role", None), "name", None)
    target_role = getattr(getattr(assigned_user, "role", None), "name", None)

    if not current_role:
        raise HTTPException(
            status_code=403,
            detail="Current user role not found"
        )

    if not target_role:
        raise HTTPException(
            status_code=404,
            detail="Assigned user role not found"
        )

    if current_role == "SuperAdmin":

        if target_role != "Admin":

            raise HTTPException(
                status_code=403,
                detail="SuperAdmin can assign task only to Admin"
            )

    elif current_role == "Admin":

        if target_role not in {"Employee", "Admin"}:

            raise HTTPException(
                status_code=403,
                detail="Admin can assign personal tasks only to Employee or Admin users"
            )

    else:

        raise HTTPException(
            status_code=403,
            detail="You are not allowed to assign tasks"
        )

    task = Task(

        title=data.title,
        description=data.description,
        priority=data.priority,
        deadline=data.deadline,

        task_type="Personal",
        group_id=None,

        created_by=current_user.id,
        assigned_by=current_user.id,
        assigned_to=assigned_user.id,
        submitted_to=current_user.id,

        is_self_task=False
    )
    db.add(task)

    notification = Notification(

        user_id=assigned_user.id,

        title="New Personal Task Assigned",

        message=f"You received task '{task.title}'",

        notification_type=f"task_assigned '{task.task_type}'"
    )
    print("Current User:", current_user.id)
    print("Assigned To:", assigned_user.id)

    db.add(notification)

    await db.commit()

    await db.refresh(task)

    # fetch assigned_by and assigned_to user records for names
    assigned_by_user = None
    assigned_to_user = None

    if task.assigned_by:
        assigned_by_result = await db.execute(select(User).where(User.id == task.assigned_by))
        assigned_by_user = assigned_by_result.scalar_one_or_none()

    if task.assigned_to:
        assigned_to_result = await db.execute(select(User).where(User.id == task.assigned_to))
        assigned_to_user = assigned_to_result.scalar_one_or_none()

    return {
        "message": "Task assigned successfully",
        "task": {
            "id": task.id,
            "title": task.title,
            "description": task.description,
            "status": task.status,
            "approval_status": task.approval_status,
            "priority": task.priority,

            "created_by": {
                "id": current_user.id,
                "username": current_user.username
            },

            "assigned_by": {
                "id": current_user.id,
                "username": current_user.username
            },

            "assigned_to": {
                "id": assigned_user.id,
                "username": assigned_user.username
            },

            "submitted_to": {
                "id": current_user.id,
                "username": current_user.username
            },

            "task_type": task.task_type,
            "group_id": task.group_id,
            "is_self_task": task.is_self_task,
            "deadline": task.deadline,
            "created_at": task.created_at
        }
    }
        
   

async def create_self_task_service(db: AsyncSession,data,current_user):

    task = Task(

        title=data.title,

        description=data.description,

        priority=data.priority,
        
        deadline = data.deadline,

        created_by=current_user.id,

        assigned_by=current_user.id,

        assigned_to=current_user.id,

        submitted_to=current_user.id,

        is_self_task=True
    )

    db.add(task)

    await db.commit()

    await db.refresh(task)

    return task


async def get_self_tasks(db:AsyncSession ,current_user):
    result= await db.execute(
        select(Task).where(
            Task.created_by == current_user.id,
            Task.is_self_task == True,
            Task.is_active == True
        )
    )
    return result.scalars().all()

async def update_task_status_service(db: AsyncSession,task_id: int,data,current_user):

    result = await db.execute(
        select(Task).where(
            Task.id == task_id
        )
    )

    task = result.scalar_one_or_none()

    if not task:
        raise HTTPException(status_code=404,detail="Task not found")

    if task.assigned_to != current_user.id:
        raise HTTPException(status_code=403,detail="Not allowed")

    task.status = data.status

    notification = Notification(

        user_id=task.submitted_to,

        title="Task Status Updated",

        message=f"Task '{task.title}' moved to {data.status}",

        notification_type="status_changed"
    )

    db.add(notification)

    await db.commit()

    await db.refresh(task)

    return task



async def add_task_log_service(db: AsyncSession,data,current_user):

    result = await db.execute(
        select(Task).where(
            Task.id == data.task_id
        )
    )

    task = result.scalar_one_or_none()

    if not task:
        raise HTTPException(status_code=404,detail="Task not found")


    # Check if this user alredy logged this task
    existing_log = await db.scalar(
        select(TaskLog).where(
            TaskLog.task_id == data.task_id,
            TaskLog.user_id == current_user.id
        )
    )

    if existing_log:
        raise HTTPException(status_code=400,detail="You have already logged this task")

    log = TaskLog(

        task_id=data.task_id,

        user_id=current_user.id,

        work_note=data.work_note,

        hours_spent=data.hours_spent
    )

    db.add(log)

    await db.commit()   

    await db.refresh(log)

    return log



async def submit_task_service(db: AsyncSession,task_id: int,current_user):

    result = await db.execute(
        select(Task).where(
            Task.id == task_id
        )
    )

    task = result.scalar_one_or_none()

    if not task:
        raise HTTPException(status_code=404,detail="Task not found")

    task.status =TaskStatus.SUBMITTED
    task.approval_status = ApprovalStatus.SUBMITTED
    # i want to in the task completed date
    task.completed_at = datetime.utcnow()

    notification = Notification(

        user_id=task.submitted_to,

        title="Task Submitted",

        message=f"Task '{task.title}' ({task.task_type}) submitted for review",

        notification_type="task_submitted"
    )

    db.add(notification)

    await db.commit()

    await db.refresh(task)

    return task



async def review_task_service(db: AsyncSession,task_id: int,data,current_user):

    result = await db.execute(
        select(Task).where(
            Task.id == task_id
        )
    )

    task = result.scalar_one_or_none()

    if not task:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    current_role = getattr(getattr(current_user, "role", None), "name", None)
    is_assigner = task.assigned_by == current_user.id
    is_group_leader = False

    if task.group_id:
        leader_result = await db.execute(
            select(GroupMember).where(
                GroupMember.group_id == task.group_id,
                GroupMember.user_id == current_user.id,
                GroupMember.role_in_group == "Lead",
                GroupMember.is_active == True
            )
        )
        is_group_leader = leader_result.scalar_one_or_none() is not None

    if current_role not in {"Admin", "SuperAdmin"} and not is_assigner and not is_group_leader:
        raise HTTPException(
            status_code=403,
            detail="Only the task assigner, group leader, admin, or superadmin can review this task"
        )

    # task.approval_status = data.review_status
    if data.review_status == ApprovalStatus.APPROVED:
        task.approval_status =ApprovalStatus.APPROVED
        task.status = TaskStatus.COMPLETED
    elif data.review_status == ApprovalStatus.NEEDS_REVISION:
        task.approval_status = ApprovalStatus.NEEDS_REVISION
        task.status = TaskStatus.IN_PROGRESS

    review = TaskReview(

        task_id=task.id,

        reviewed_by=current_user.id,

        reviewed_status=data.review_status,

        comment=data.comment
    )

    db.add(review)

    notification = Notification(

        user_id=task.assigned_to,

        title="Task Review Update",

        message=data.comment,

        notification_type=data.review_status
    )

    db.add(notification)

    await db.commit()

    return {
        "message": "Task reviewed successfully"
    }


async def upload_task_evidence_service(db: AsyncSession,task_id: int,description: str,current_user,screenshot=None,file=None):

    result = await db.execute(
        select(Task).where(Task.id == task_id)
    )

    task = result.scalar_one_or_none()

    if not task:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    current_role = getattr(getattr(current_user, "role", None), "name", None)
    is_assigner = task.assigned_by == current_user.id
    is_assigned_employee = task.assigned_to == current_user.id
    is_group_leader = False

    if task.group_id:
        leader_result = await db.execute(
            select(GroupMember).where(
                GroupMember.group_id == task.group_id,
                GroupMember.user_id == current_user.id,
                GroupMember.role_in_group == "Lead",
                GroupMember.is_active == True
            )
        )
        is_group_leader = leader_result.scalar_one_or_none() is not None

    if current_role not in {"Admin", "SuperAdmin"} and not is_assigner and not is_group_leader and not is_assigned_employee:
        raise HTTPException(
            status_code=403,
            detail="Only the assigned employee, task assigner, group leader, admin, or superadmin can upload evidence"
        )

    screenshot_path = None
    file_path = None

    if screenshot is not None:
        screenshot_path = await save_upload_file(screenshot, "task_evidences/screenshots")

    if file is not None:
        file_path = await save_upload_file(file, "task_evidences/files")

    evidence = TaskEvidence(
        task_id=task_id,
        user_id=current_user.id,
        description=description,
        screenshot_path=screenshot_path,
        file_path=file_path
    )

    db.add(evidence)
    await db.commit()
    await db.refresh(evidence)

    return evidence


async def get_task_evidence_service(db: AsyncSession, task_id: int, current_user):

    task_result = await db.execute(
        select(Task).where(Task.id == task_id)
    )

    task = task_result.scalar_one_or_none()

    if not task:
        raise HTTPException(status_code=404, detail="Task not found")

    current_role = getattr(getattr(current_user, "role", None), "name", None)
    evidence_query = (
        select(TaskEvidence)
        .where(TaskEvidence.task_id == task_id)
        .order_by(TaskEvidence.created_at.desc())
    )

    if current_role in {"Admin", "SuperAdmin"}:
        result = await db.execute(evidence_query)
        return result.scalars().all()

    if task.group_id:
        leader_result = await db.execute(
            select(GroupMember).where(
                GroupMember.group_id == task.group_id,
                GroupMember.user_id == current_user.id,
                GroupMember.role_in_group == "Lead",
                GroupMember.is_active == True
            )
        )
        if leader_result.scalar_one_or_none():
            result = await db.execute(evidence_query)
            return result.scalars().all()

    if task.assigned_to == current_user.id or task.assigned_by == current_user.id:
        result = await db.execute(evidence_query)
        return result.scalars().all()

    raise HTTPException(
        status_code=403,
        detail="You do not have permission to view this task evidence"
    )


async def get_personal_assigned_tasks(db: AsyncSession, current_user):

    result = await db.execute(
        select(
            Task.id,
            Task.title,
            Task.description,
            Task.priority,
            Task.status,
            Task.approval_status,
            Task.deadline,
            Task.created_at,
            User.username.label("assigned_by_name")
        )
        .join(
            User,
            User.id == Task.assigned_by      # Join with assigner
        )
        .where(
            Task.assigned_to == current_user.id,
            Task.task_type == "Personal",
            Task.is_self_task.is_(False)
        )
        .order_by(Task.created_at.desc())
    )

    return result.mappings().all()
async def get_my_tasks_service(db: AsyncSession, current_user):

    result = await db.execute(
        select(Task)
        .where(
            Task.assigned_to == current_user.id,
            Task.task_type == "Personal",
            Task.assigned_by == current_user.id,
            Task.is_self_task == False,
            Task.is_active == True
        )
        .order_by(Task.created_at.desc(), Task.id.desc())
    )

    return result.scalars().all()

async def get_group_tasks_service(
    db: AsyncSession,
    current_user,
    group_id: int
):

    AssignedByUser = aliased(User)
    AssignedToUser = aliased(User)

    result = await db.execute(
        select(
            Task,
            AssignedByUser.username.label("assigned_by_name"),
            AssignedToUser.username.label("assigned_to_name")
        )
        .outerjoin(
            AssignedByUser,
            Task.assigned_by == AssignedByUser.id
        )
        .outerjoin(
            AssignedToUser,
            Task.assigned_to == AssignedToUser.id
        )
        .where(
    Task.group_id == group_id,
    Task.task_type == "Group",
    Task.is_active == True
)
        .order_by(
            Task.created_at.desc(),
            Task.id.desc()
        )
    )

    rows = result.all()

    response = []

    for task, assigned_by_name, assigned_to_name in rows:

        task_data = {
            column.name: getattr(task, column.name)
            for column in Task.__table__.columns
        }

        task_data["assigned_by_name"] = assigned_by_name
        task_data["assigned_to_name"] = assigned_to_name

        response.append(task_data)

    return response



async def get_all_tasks_service(
    db: AsyncSession,
    group_id: int,
    current_user
):


    group_member_result = await db.execute(
        select(GroupMember).where(
            GroupMember.group_id == group_id,
            GroupMember.user_id == current_user.id
        )
    )

    group_member = group_member_result.scalar_one_or_none()

    if not group_member:
        raise HTTPException(
            status_code=403,
            detail="You are not member of this group"
        )


    if group_member.role_in_group == "Lead":

        result = await db.execute(
            select(Task)
            .where(
                Task.group_id == group_id,
                Task.is_active == True
            )
            .order_by(Task.created_at.desc(), Task.id.desc())
        )

        return result.scalars().all()


    result = await db.execute(
        select(Task)
        .where(
            Task.group_id == group_id,
            Task.assigned_to == current_user.id,
            Task.is_active == True
        )
        .order_by(Task.created_at.desc(), Task.id.desc())
    )

    return result.scalars().all()



async def get_task_by_id_service(db,task_id:int,current_user):
    result = await db.execute(select(Task).where(Task.id==task_id,Task.is_active==True))
    task = result.scalar_one_or_none()

    if not task:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    # Allow only:
    # 1. Assigned Employee
    # 2. Task Creator / Group Leader
    # 3. Admin / SuperAdmin
    print("Current User ID:", current_user.id)
    print("Current User Role:", current_user.role)
    print("Task Assigned To:", task.assigned_to)
    print("Task Assigned By:", task.assigned_by)
    if (
        current_user.role.name not in ["Admin", "SuperAdmin"]
        and task.assigned_to != current_user.id
        and task.assigned_by != current_user.id
    ):
        raise HTTPException(
            status_code=403,
            detail="You are not authorized to view this task"
        )

    # fetch assigned_by and assigned_to user records for names
    assigned_by_user = None
    assigned_to_user = None

    if task.assigned_by:
        assigned_by_result = await db.execute(select(User).where(User.id == task.assigned_by))
        assigned_by_user = assigned_by_result.scalar_one_or_none()

    if task.assigned_to:
        assigned_to_result = await db.execute(select(User).where(User.id == task.assigned_to))
        assigned_to_user = assigned_to_result.scalar_one_or_none()

    print("Current User ID:", current_user.id)
    print("Current User Role:", current_user.role)
    print("Task Assigned To:", task.assigned_to)
    print("Task Assigned By:", task.assigned_by)
    print("Authorization Condition:",
      current_user.role not in ["Admin", "SuperAdmin"],
      task.assigned_to != current_user.id,
      task.assigned_by != current_user.id)
    return {
        "id": task.id,
        "title": task.title,
        "description": task.description,
        "status": task.status,
        "approval_status": task.approval_status,
        "priority": task.priority,
        "created_by": task.created_by,
        "assigned_by": task.assigned_by,
        "assigned_to": task.assigned_to,
        "task_type": task.task_type,
        "submitted_to": task.submitted_to,
        "group_id": task.group_id,
        "deadline": task.deadline,
        "created_at": task.created_at,
        "assigned_by_name": assigned_by_user.username if assigned_by_user else None,
        "assigned_to_name": assigned_to_user.username if assigned_to_user else None
    }
async def get_employee_task_review(employee_id: int,db: AsyncSession):

    result = await db.execute(

        select(
            TaskReview.id,
            TaskReview.task_id,
            Task.title.label("task_title"),
            User.username.label("reviewer_name"),
            TaskReview.reviewed_status,
            TaskReview.comment,
            TaskReview.created_at
        )

        .join(
            Task,
            Task.id == TaskReview.task_id
        )

        .join(
            User,
            User.id == TaskReview.reviewed_by
        )

        .where(
            Task.assigned_to == employee_id
        )

        .order_by(
            TaskReview.created_at.desc(),
            TaskReview.id.desc()
        )

    )

    reviews = result.mappings().all()

    if not reviews:
        raise HTTPException(
            status_code=404,
            detail="No task reviews found"
        )

    return reviews


async def  get_task_review_by_id(taskId:int ,db:AsyncSession):
    last_24_hours = datetime.now() - timedelta(hours=24)    
    result = await db.execute(select(
        TaskReview.id,
        TaskReview.task_id,
        Task.title.label("task_title"),
        User.username.label("reviewer_name"),
        TaskReview.reviewed_status,
        TaskReview.comment,
        TaskReview.created_at
    ).join(
        Task,
        Task.id == TaskReview.task_id
    ).join(
        User,
        User.id == TaskReview.reviewed_by
        
    )
    .where(
        TaskReview.task_id ==  taskId,
        TaskReview.created_at >= last_24_hours
    ).order_by(
        TaskReview.created_at.desc(),
        TaskReview.id.desc()
        
    )
                            
)
    reviews =result.mappings().all()
    
    if not reviews:
        raise HTTPException(status_code=404, detail="Reviews Not Found")

    return reviews

async def soft_delete_group_task_service(db: AsyncSession,task_id: int,current_user):


    task_result = await db.execute(
        select(Task).where(
            Task.id == task_id,
            Task.is_active == True
        )
    )
    task = task_result.scalar_one_or_none()

    if not task:
        raise HTTPException(
            status_code=404,
            detail="Task not found"
        )

    if not task.group_id:
        raise HTTPException(
            status_code=400,
            detail="Only group tasks can be soft deleted by a group leader"
        )

    leader_result = await db.execute(
        select(GroupMember).where(
            GroupMember.group_id == task.group_id,
            GroupMember.user_id == current_user.id,
            GroupMember.role_in_group == "Lead",
            GroupMember.is_active == True
        )
    )
    leader = leader_result.scalar_one_or_none()

    if not leader:
        raise HTTPException(
            status_code=403,
            detail="Only the group leader of this task can soft delete it"
        )

    task.is_active = False

    await db.commit()

    return {
        "message": "Task soft deleted successfully",
        "task_id": task.id,
        "is_active": task.is_active
    }


async def get_personal_tasks_for_review_service(db: AsyncSession, current_user):

    role = getattr(getattr(current_user, "role", None), "name", None)

    base_query = (
        select(Task)
        .where(
            Task.task_type == "Personal",
            Task.is_self_task == False,
            Task.approval_status == ApprovalStatus.SUBMITTED
        )
        .order_by(Task.created_at.desc())
    )

    # 👑 SuperAdmin → sees ONLY tasks assigned BY SuperAdmin
    if role == "SuperAdmin":

        result = await db.execute(
            base_query.where(
                Task.assigned_by == current_user.id
            )
        )
        return result.scalars().all()

    # 🧑‍💼 Admin → sees ONLY tasks assigned BY Admin
    if role == "Admin":

        result = await db.execute(
            base_query.where(
                Task.assigned_by == current_user.id
            )
        )
        return result.scalars().all()

    raise HTTPException(
        status_code=403,
        detail="Not allowed"
    )

async def export_task_report_service(
    db: AsyncSession,
    current_user,
    report_type: str,
    selected_date: datetime = None
):

    now = datetime.utcnow()

    def _to_naive_utc(dt: datetime):
        if dt is None:
            return None

        if dt.tzinfo is not None:
            return dt.astimezone(timezone.utc).replace(
                tzinfo=None
            )

        return dt

    if report_type == "daily":

        start_date = now.replace(
            hour=0,
            minute=0,
            second=0,
            microsecond=0
        )

    elif report_type == "weekly":

        start_date = now - timedelta(days=7)

    elif report_type == "monthly":

        start_date = now - timedelta(days=30)

    elif report_type == "yearly":

        start_date = now - timedelta(days=365)

    else:

        raise HTTPException(
            status_code=400,
            detail="Invalid report type"
        )

    if selected_date:

        sel = _to_naive_utc(selected_date)

        start_date = sel.replace(
            hour=0,
            minute=0,
            second=0,
            microsecond=0
        )

        end_date = sel.replace(
            hour=23,
            minute=59,
            second=59,
            microsecond=999999
        )

        result = await db.execute(
            select(TaskLog, Task)
            .join(
                Task,
                Task.id == TaskLog.task_id
            )
            .where(
                TaskLog.user_id == current_user.id,
                TaskLog.created_at >= start_date,
                TaskLog.created_at <= end_date
            )
        )

    else:

        result = await db.execute(
            select(TaskLog, Task)
            .join(
                Task,
                Task.id == TaskLog.task_id
            )
            .where(
                TaskLog.user_id == current_user.id,
                TaskLog.created_at >= start_date
            )
        )

    logs = result.all()

    if not logs:

        raise HTTPException(
            status_code=404,
            detail="No logs found"
        )

    report_data = []

    total_hours_worked = 0
    total_extra_hours = 0
    total_leaves = 0

    for log, task in logs:

        worked_hours = float(
            log.hours_spent or 0
        )

        extra_hours = max(
            worked_hours - 8,
            0
        )
        print(f'the extra hours value is this ${extra_hours}')
        total_hours_worked += worked_hours
        total_extra_hours += extra_hours

        report_data.append({

            "Sr No": len(report_data) + 1,

            "Date": log.created_at.strftime(
                "%d-%m-%Y"
            ),

            "Day": log.created_at.strftime(
                "%A"
            ),

            "Task Title": task.title,

            "Project Name": (
                task.project_name
                if task.project_name
                else "N/A"
            ),

            "Work Description": (
                log.work_note
            ),

            "Hours Worked": worked_hours,

            "Extra Hours": extra_hours,

            "Leaves": 0

        })

    reports_dir = os.path.join(
        "uploads",
        "reports"
    )

    os.makedirs(
        reports_dir,
        exist_ok=True
    )

    file_name = (
        f"{report_type}_report_"
        f"{current_user.id}.xlsx"
    )

    file_path = os.path.join(
        reports_dir,
        file_name
    )

    headers = [

        "Sr No",

        "Date",

        "Day",

        "Task Title",

        "Project Name",

        "Work Description",

        "Hours Worked",

        "Extra Hours",

        "Leaves"

    ]

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Task Report"

    sheet.append(headers)

    for row in report_data:
        sheet.append([
            row["Sr No"],
            row["Date"],
            row["Day"],
            row["Task Title"],
            row["Project Name"],
            row["Work Description"],
            row["Hours Worked"],
            row["Extra Hours"],
            row["Leaves"],
        ])

    sheet.append([])
    sheet.append(["WORK SUMMARY", "", "", "", "", "", "", "", ""])
    sheet.append(["Total Tasks Logged", len(logs), "", "", "", "", "", "", ""])
    sheet.append(["Total Hours Worked", total_hours_worked, "", "", "", "", "", "", ""])
    sheet.append(["Total Extra Hours", total_extra_hours, "", "", "", "", "", "", ""])
    sheet.append(["Total Leaves Taken", total_leaves, "", "", "", "", "", "", ""])

    header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    bold_font = Font(bold=True)

    for cell in sheet[1]:
        cell.font = bold_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    for column_cells in sheet.columns:
        max_length = 0
        for cell in column_cells:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        adjusted_width = min(max_length + 2, 50)
        sheet.column_dimensions[column_cells[0].column_letter].width = adjusted_width

    workbook.save(file_path)

    return FileResponse(
        path=file_path,
        filename=file_name,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    
    
    